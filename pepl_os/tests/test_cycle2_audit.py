"""C2-01 - Live audit workbook."""

from io import BytesIO

import frappe

from pepl_os.pepl_stores import audit
from pepl_os.tests.factories import make_item
from pepl_os.tests.utils import PEPLTestCase, make_user


class TestCycle2Audit(PEPLTestCase):
	def test_every_sheet_is_built(self):
		result = audit.build_audit("2026-07-01")
		self.assertEqual(list(result), [title for title, _cols in audit.SHEETS])
		for rows in result.values():
			self.assertIsInstance(rows, list)
		self.assertEqual(len(result["7 Vendor bills"]), 1)

	def test_item_outside_tree_is_listed(self):
		frappe.flags.pepl_allow_items_outside_tree = True
		try:
			make_item("PEPL-T-AUDIT-OUT", "_Test Item Group")
		finally:
			frappe.flags.pepl_allow_items_outside_tree = False
		rows = audit.build_audit("2026-07-01")["3 Items"]
		self.assertIn(
			"PEPL-T-AUDIT-OUT", {r["item"] for r in rows if r["check"].startswith("Item Group outside")}
		)

	def test_workbook_has_summary_and_sheets(self):
		from openpyxl import load_workbook

		content = audit.to_xlsx(audit.build_audit("2026-07-01"))
		wb = load_workbook(BytesIO(content))
		self.assertEqual(wb.sheetnames, ["Summary", *[title for title, _cols in audit.SHEETS]])
		self.assertEqual(wb["2 Suppliers"].cell(1, 4).value, "Owner")

	def test_stock_value_line_counts_items(self):
		# Regression (found on the demo site): the SQL alias "items" clashed with dict.items(),
		# so the line read "across <built-in method items ...> item(s)".
		for row in audit.build_audit("2026-07-01")["4 Stock"]:
			if row["check"] == "Stock value by warehouse":
				self.assertNotIn("built-in", row["detail"])
				self.assertRegex(row["detail"], r"across \d+ item\(s\)")

	def test_summary_counts_only_findings(self):
		from openpyxl import load_workbook

		data = {title: [] for title, _cols in audit.SHEETS}
		data["7 Vendor bills"] = [{"check": "Vendor bills in ERPNext", "detail": "info"}]
		data["2 Suppliers"] = [{"check": "Missing GSTIN", "supplier": "X", "detail": "X"}]
		summary = load_workbook(BytesIO(audit.to_xlsx(data)))["Summary"]
		counts = {summary.cell(r, 1).value: summary.cell(r, 2).value for r in range(2, summary.max_row + 1)}
		self.assertEqual((counts["7 Vendor bills"], counts["2 Suppliers"]), (0, 1))

	def test_audit_writes_nothing(self):
		before = frappe.db.count("File")
		audit.to_xlsx(audit.build_audit("2026-07-01"))
		self.assertEqual(frappe.db.count("File"), before)

	def test_download_is_system_manager_only(self):
		user = make_user("pepl.audit.stores@example.com")
		frappe.set_user(user)
		try:
			with self.assertRaises(frappe.PermissionError):
				audit.download_cycle2_audit()
		finally:
			frappe.set_user("Administrator")
