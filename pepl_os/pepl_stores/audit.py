"""C2-01 - Read-only live audit of purchase and stores data before Cycle 2 goes live.

Run by the System Manager from PEPL System Health ("Download Cycle 2 audit").
Nothing is written to the database: the workbook is built in memory and
downloaded. One sheet per check, each finding with empty Owner and Fix
columns to fill in.
"""

from io import BytesIO

import frappe
from frappe.utils import add_days, cstr, flt, getdate, today

from pepl_os.pepl_stores import stock_tree

DEFAULT_SINCE = "2026-07-01"
STALE_DRAFT_DAYS = 7

# Buying and stock DocTypes whose UI customisations must be known before Cycle 2 adds its own.
BUY_STOCK_DOCTYPES = (
	"Supplier",
	"Item",
	"Item Group",
	"Warehouse",
	"Material Request",
	"Request for Quotation",
	"Supplier Quotation",
	"Purchase Order",
	"Purchase Receipt",
	"Purchase Invoice",
	"Stock Entry",
	"Stock Reconciliation",
	"Delivery Note",
	"Quality Inspection",
	"Batch",
)
BUY_STOCK_MODULES = ("Buying", "Stock")


def _field(doctype, *candidates):
	meta = frappe.get_meta(doctype)
	return next((f for f in candidates if meta.has_field(f)), None)


def _purchase_documents(since):
	rows = []
	stale_before = add_days(today(), -STALE_DRAFT_DAYS)
	date_field = {
		"Purchase Order": "transaction_date",
		"Purchase Receipt": "posting_date",
		"Purchase Invoice": "posting_date",
	}
	for doctype, dfield in date_field.items():
		base = {dfield: [">=", since]}
		counts = {
			label: frappe.db.count(doctype, {**base, "docstatus": status})
			for label, status in (("Draft", 0), ("Submitted", 1), ("Cancelled", 2))
		}
		rows.append({"check": "Count", "doctype": doctype, "document": "", "detail": str(counts)})
		for d in frappe.get_all(
			doctype,
			filters={**base, "docstatus": 0, "creation": ["<", stale_before]},
			fields=["name", "creation", "owner"],
			order_by="creation asc",
		):
			rows.append(
				{
					"check": f"Draft older than {STALE_DRAFT_DAYS} days",
					"doctype": doctype,
					"document": d.name,
					"detail": f"Created {getdate(d.creation)} by {d.owner}",
				}
			)
	for line in frappe.get_all(
		"Purchase Order Item",
		filters={"schedule_date": ["is", "not set"], "docstatus": ["<", 2]},
		fields=["parent", "idx", "item_code"],
		order_by="parent asc, idx asc",
	):
		rows.append(
			{
				"check": "PO line without delivery date",
				"doctype": "Purchase Order",
				"document": line.parent,
				"detail": f"Row {line.idx}: {line.item_code}",
			}
		)
	return rows


def _suppliers():
	rows = []
	gstin = _field("Supplier", "gstin", "tax_id")
	pan = _field("Supplier", "pan")
	fields = ["name", "supplier_name", "email_id"] + [f for f in (gstin, pan) if f]
	suppliers = frappe.get_all("Supplier", filters={"disabled": 0}, fields=fields)
	for key, label in ((gstin, "GSTIN"), (pan, "PAN")):
		if not key:
			continue
		seen = {}
		for s in suppliers:
			value = cstr(s.get(key)).strip().upper()
			if value:
				seen.setdefault(value, []).append(s.name)
		for value, names in seen.items():
			if len(names) > 1:
				rows.append({"check": f"Duplicate {label}", "supplier": ", ".join(names), "detail": value})
	for s in suppliers:
		if gstin and not s.get(gstin):
			rows.append({"check": "Missing GSTIN", "supplier": s.name, "detail": s.supplier_name})
		if not s.email_id:
			rows.append(
				{
					"check": "Missing email (needed for chase letters)",
					"supplier": s.name,
					"detail": s.supplier_name,
				}
			)
	return rows


def _has_default_warehouse(item, company):
	if not company:
		return True
	if frappe.db.exists(
		"Item Default", {"parent": item.name, "company": company, "default_warehouse": ["is", "set"]}
	):
		return True
	return bool(
		frappe.db.exists(
			"Item Default",
			{
				"parent": item.item_group,
				"parenttype": "Item Group",
				"company": company,
				"default_warehouse": ["is", "set"],
			},
		)
	)


def _items(company):
	rows = []
	for item in frappe.get_all(
		"Item",
		filters={"disabled": 0, "is_stock_item": 1},
		fields=["name", "item_name", "item_group", "valuation_rate"],
	):
		stock_class = stock_tree.stock_class_of(item.item_group)
		if not stock_class:
			rows.append(
				{
					"check": "Item Group outside the stock classes",
					"item": item.name,
					"item_group": item.item_group,
					"detail": item.item_name,
				}
			)
		elif stock_class == stock_tree.CONSUMABLES and not flt(item.valuation_rate):
			rows.append(
				{
					"check": "Consumable with no valuation rate",
					"item": item.name,
					"item_group": item.item_group,
					"detail": item.item_name,
				}
			)
		if not _has_default_warehouse(item, company):
			rows.append(
				{
					"check": "No default warehouse",
					"item": item.name,
					"item_group": item.item_group,
					"detail": item.item_name,
				}
			)
	return rows


def _stock():
	rows = []
	for w in frappe.db.sql(
		"""select warehouse, sum(stock_value) as value, count(*) as items
		from `tabBin` group by warehouse order by value desc""",
		as_dict=True,
	):
		rows.append(
			{
				"check": "Stock value by warehouse",
				"warehouse": w.warehouse,
				"item": "",
				"detail": f"{flt(w.value, 2)} across {w.items} item(s)",
			}
		)
	for b in frappe.get_all(
		"Bin", filters={"actual_qty": ["<", 0]}, fields=["item_code", "warehouse", "actual_qty"]
	):
		rows.append(
			{
				"check": "Negative stock",
				"warehouse": b.warehouse,
				"item": b.item_code,
				"detail": str(b.actual_qty),
			}
		)
	for b in frappe.get_all("Bin", filters={"actual_qty": [">", 0]}, fields=["item_code", "warehouse"]):
		group = frappe.get_cached_value("Item", b.item_code, "item_group")
		stock_class = stock_tree.stock_class_of(group)
		csm_item = stock_class == stock_tree.CSM or frappe.get_cached_value(
			"Item", b.item_code, "is_customer_provided_item"
		)
		csm_store = stock_tree.is_csm_warehouse(b.warehouse)
		store_name = frappe.get_cached_value("Warehouse", b.warehouse, "warehouse_name")
		expected = stock_tree.default_store_for(stock_class) if stock_class else None
		if bool(csm_item) != csm_store:
			detail = "Customer material outside CSM stores" if csm_item else "PEPL material in a CSM store"
		elif expected and store_name in stock_tree.LEAF_STORES and store_name != expected:
			detail = f"{stock_class} item in {store_name} (expected {expected})"
		else:
			continue
		rows.append(
			{
				"check": "Stock in a store of the wrong class",
				"warehouse": b.warehouse,
				"item": b.item_code,
				"detail": detail,
			}
		)
	return rows


def _heat_readiness():
	rows = []
	for item in frappe.get_all(
		"Item", filters={"disabled": 0, "is_stock_item": 1}, fields=["name", "item_group", "has_batch_no"]
	):
		stock_class = stock_tree.stock_class_of(item.item_group)
		if stock_class not in (stock_tree.RAW_MATERIAL, stock_tree.BOUGHT_OUT):
			continue
		entries = frappe.db.count("Stock Ledger Entry", {"item_code": item.name, "is_cancelled": 0})
		if entries:
			rows.append(
				{
					"check": "RM / bought-out item already has stock history",
					"item": item.name,
					"detail": f"{entries} ledger entries; batch tracked: {'Yes' if item.has_batch_no else 'No'}",
				}
			)
	return rows


def _customisations():
	rows = []
	for s in frappe.get_all("Server Script", fields=["name", "script_type", "reference_doctype", "disabled"]):
		if s.script_type != "DocType Event" or s.reference_doctype in BUY_STOCK_DOCTYPES:
			rows.append(
				{
					"check": "Server Script",
					"name": s.name,
					"doctype": s.reference_doctype or s.script_type,
					"detail": "Disabled" if s.disabled else "Enabled",
				}
			)
	for c in frappe.get_all(
		"Client Script", filters={"dt": ["in", BUY_STOCK_DOCTYPES]}, fields=["name", "dt", "enabled"]
	):
		rows.append(
			{
				"check": "Client Script",
				"name": c.name,
				"doctype": c.dt,
				"detail": "Enabled" if c.enabled else "Disabled",
			}
		)
	for d in frappe.get_all("DocType", filters={"custom": 1}, fields=["name", "module"]):
		rows.append(
			{"check": "Custom DocType (built in the UI)", "name": d.name, "doctype": d.module, "detail": ""}
		)
	for r in frappe.get_all(
		"Report", filters={"is_standard": "No"}, fields=["name", "ref_doctype", "report_type"]
	):
		module = frappe.get_cached_value("DocType", r.ref_doctype, "module") if r.ref_doctype else None
		if r.ref_doctype in BUY_STOCK_DOCTYPES or module in BUY_STOCK_MODULES:
			rows.append(
				{
					"check": "Non-standard Report",
					"name": r.name,
					"doctype": r.ref_doctype,
					"detail": r.report_type,
				}
			)
	return rows


def _vendor_bills(since):
	submitted = frappe.db.count("Purchase Invoice", {"posting_date": [">=", since], "docstatus": 1})
	receipts = frappe.db.count("Purchase Receipt", {"posting_date": [">=", since], "docstatus": 1})
	verdict = (
		"Vendor bills are entered in ERPNext"
		if submitted
		else "No vendor bills in ERPNext: bills are kept elsewhere (Tally)"
	)
	return [
		{
			"check": "Vendor bills in ERPNext",
			"detail": f"{submitted} Purchase Invoices and {receipts} Purchase Receipts submitted since {since}. {verdict}.",
		}
	]


SHEETS = (
	("1 Purchase documents", ("check", "doctype", "document", "detail")),
	("2 Suppliers", ("check", "supplier", "detail")),
	("3 Items", ("check", "item", "item_group", "detail")),
	("4 Stock", ("check", "warehouse", "item", "detail")),
	("5 Heat readiness", ("check", "item", "detail")),
	("6 UI customisations", ("check", "name", "doctype", "detail")),
	("7 Vendor bills", ("check", "detail")),
)


def build_audit(since=DEFAULT_SINCE, company=None):
	"""{sheet title: [row dicts]} for every check. Reads only."""
	since = getdate(since)
	company = company or stock_tree.default_company()
	data = [
		_purchase_documents(since),
		_suppliers(),
		_items(company),
		_stock(),
		_heat_readiness(),
		_customisations(),
		_vendor_bills(since),
	]
	return {title: rows for (title, _cols), rows in zip(SHEETS, data, strict=True)}


def to_xlsx(audit, sheets=None):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	wb = Workbook()
	summary = wb.active
	summary.title = "Summary"
	summary.append(["Sheet", "Findings"])
	header_font = Font(bold=True, color="FFFFFF")
	header_fill = PatternFill("solid", start_color="1F3A5F")
	for title, columns in sheets or SHEETS:
		rows = audit.get(title, [])
		summary.append([title, len(rows)])
		ws = wb.create_sheet(title[:31])
		ws.append([c.replace("_", " ").title() for c in columns] + ["Owner", "Fix"])
		for row in rows:
			ws.append([cstr(row.get(c)) for c in columns] + ["", ""])
		for cell in ws[1]:
			cell.font, cell.fill = header_font, header_fill
		for i, width in enumerate([34] + [28] * (len(columns) - 1) + [18, 30], start=1):
			ws.column_dimensions[ws.cell(1, i).column_letter].width = width
	for cell in summary[1]:
		cell.font, cell.fill = header_font, header_fill
	summary.column_dimensions["A"].width = 28
	out = BytesIO()
	wb.save(out)
	return out.getvalue()


@frappe.whitelist()
def download_cycle2_audit(since=DEFAULT_SINCE):
	"""System Manager only. Streams the audit workbook; writes nothing."""
	frappe.only_for("System Manager")
	content = to_xlsx(build_audit(since))
	frappe.response["filename"] = f"PEPL Cycle 2 audit {today()}.xlsx"
	frappe.response["filecontent"] = content
	frappe.response["type"] = "binary"
	return None
